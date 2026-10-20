"""One-off: add a "Shopping & misc" budget line to Sabbatical_Plan_and_Budget.xlsx.

- Itinerary: new column O "Shopping & misc (£)" per stop (default £5 a night), total in O25,
  and each stop's Total (column K) now includes it.
- Overview: new "Shopping & misc" row in "2. Budget by category" (row 25); the rows below move
  down one and their formulas are adjusted.
- Tracker: "Shopping & misc" added to the Category drop-down.

Safe to run twice (does nothing if the column is already there). Run LibreOffice afterwards to
recalculate (see .github/workflows/sync-tracker.yml).
Usage: python scripts/add_shopping_misc.py [xlsx] [pounds_per_night]
"""
import copy
import re
import sys

import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else "Sabbatical_Plan_and_Budget.xlsx"
PER_NIGHT = float(sys.argv[2]) if len(sys.argv) > 2 else 5
LABEL = "Shopping & misc"

wb = openpyxl.load_workbook(XLSX)
it, ov, tr = wb["Itinerary"], wb["Overview"], wb["Tracker"]
if it["O4"].value == f"{LABEL} (£)":
    print("Already done")
    sys.exit(0)

def style(dst, src):
    dst._style = copy.copy(src._style)

# ---- Itinerary
it["O4"].value = f"{LABEL} (£)"; style(it["O4"], it["H4"])
total_row = next(r for r in range(5, 200) if str(it[f"A{r}"].value or "").lower().startswith("total"))
for r in range(5, total_row):
    if not it[f"A{r}"].value:
        continue
    it[f"O{r}"].value = round(PER_NIGHT * float(it[f"D{r}"].value or 0), 2)
    style(it[f"O{r}"], it[f"H{r}"])
    it[f"K{r}"].value = f"=G{r}+H{r}+I{r}+J{r}+O{r}"
it[f"O{total_row}"].value = f"=SUM(O5:O{total_row-1})"; style(it[f"O{total_row}"], it[f"H{total_row}"])
it.column_dimensions["O"].width = 14

# ---- Overview: open up row 25 (between Activities and Insurance & admin)
AT = 25
assert ov[f"A{AT-1}"].value == "Activities" and ov[f"A{AT}"].value == "Insurance & admin", "Overview layout changed"
REF = re.compile(r"((?:'[^']+'|[A-Za-z_][\w ]*)!)?(\$?)([A-Z]{1,3})(\$?)(\d+)(?::(\$?)([A-Z]{1,3})(\$?)(\d+))?")
def shift(f):
    def one(m):
        if m.group(1):  # another sheet – leave alone
            return m.group(0)
        bump = lambda n: str(int(n) + 1 if int(n) >= AT else int(n))
        s = f"{m.group(2)}{m.group(3)}{m.group(4)}{bump(m.group(5))}"
        if m.group(7):
            s += f":{m.group(6)}{m.group(7)}{m.group(8)}{bump(m.group(9))}"
        return s
    out, parts = [], re.split(r'("[^"]*")', f)   # don't touch text inside quotes
    for p in parts:
        out.append(p if p.startswith('"') else REF.sub(one, p))
    return "".join(out)

maxr, maxc = ov.max_row, ov.max_column
for r in range(maxr, AT - 1, -1):
    for c in range(1, maxc + 1):
        src, dst = ov.cell(r, c), ov.cell(r + 1, c)
        dst.value = src.value; style(dst, src)
    if r in ov.row_dimensions and ov.row_dimensions[r].height:
        ov.row_dimensions[r + 1].height = ov.row_dimensions[r].height
for row in ov.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.value = shift(c.value)
for c in range(1, maxc + 1):
    style(ov.cell(AT, c), ov.cell(AT - 1, c)); ov.cell(AT, c).value = None
ov[f"A{AT}"].value = LABEL
ov[f"B{AT}"].value = f"=Itinerary!O{total_row}"
ov[f"C{AT}"].value = f"=SUMIFS(Tracker!$E$5:$E$1000,Tracker!$C$5:$C$1000,A{AT})"
ov[f"D{AT}"].value = f"=B{AT}-C{AT}"
ov[f"E{AT}"].value = f"=IF($B$3=0,0,B{AT}/$B$3)"
ov.row_dimensions[AT].height = ov.row_dimensions[AT - 1].height

for row in ov.iter_rows(max_row=AT - 2):
    for c in row:
        if isinstance(c.value, str) and "local transport, activities)" in c.value:
            c.value = c.value.replace("local transport, activities)", "local transport, activities, shopping & misc)")

# ---- Tracker category drop-down
for dv in tr.data_validations.dataValidation:
    if dv.formula1 and "Local transport" in dv.formula1 and LABEL not in dv.formula1:
        dv.formula1 = dv.formula1.replace("Activities,", f"Activities,{LABEL},")

wb.calculation.fullCalcOnLoad = True
wb.save(XLSX)
print(f"Added {LABEL}: £{PER_NIGHT:g}/night")
