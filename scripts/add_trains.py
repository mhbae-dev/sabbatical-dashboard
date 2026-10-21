"""One-off: add a "Trains & intercity" budget line to Sabbatical_Plan_and_Budget.xlsx.

Trains, long-distance buses and airport trains between places used to sit inside each stop's
Local transport. They now get their own line, outside the daily budget (like Flights):
- Itinerary: new column P "Trains & intercity (£)". The legs the notes said were included in Local
  transport are moved out of column G into P (rough fares), so the trip total doesn't change.
  Each stop's Total (column K) now includes P.
- Overview: new "Trains & intercity" row in "2. Budget by category" right after Local transport;
  the rows below move down one and their formulas are adjusted.
- Tracker: "Trains & intercity" added to the Category drop-down.

Safe to run twice. Recalculate in LibreOffice afterwards (see .github/workflows/sync-tracker.yml).
Usage: python scripts/add_trains.py [xlsx]
"""
import copy
import re
import sys

import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else "Sabbatical_Plan_and_Budget.xlsx"
LABEL = "Trains & intercity"
# where (start of Itinerary column B) -> (£ moved from Local transport, what it covers). Rough fares.
MOVE = {
    "Yokohama": (23, "Narita Express from Narita"),
    "Nagoya": (26, "Shizuoka → Nagoya shinkansen"),
    "Takayama": (31, "Nagoya → Takayama limited express"),
    "Tainan": (33, "Taipei → Tainan high-speed rail"),
    "Tokyo": (118, "Hakodate → Tokyo shinkansen"),
}

wb = openpyxl.load_workbook(XLSX)
it, ov, tr = wb["Itinerary"], wb["Overview"], wb["Tracker"]
if it["P4"].value == f"{LABEL} (£)":
    print("Already done"); sys.exit(0)

def style(dst, src):
    dst._style = copy.copy(src._style)

# ---- Itinerary
it["P4"].value = f"{LABEL} (£)"; style(it["P4"], it["G4"])
total_row = next(r for r in range(5, 200) if str(it[f"A{r}"].value or "").lower().startswith("total"))
for r in range(5, total_row):
    if not it[f"A{r}"].value:
        continue
    where = str(it[f"B{r}"].value or "")
    amt, what = next(((a, w) for k, (a, w) in MOVE.items() if where.startswith(k)), (0, None))
    if amt:
        it[f"G{r}"].value = float(it[f"G{r}"].value or 0) - amt
        it[f"N{r}"].value = re.sub(r"Transport incl\. [^.]*\.?\s*", "", str(it[f"N{r}"].value or "")).strip() + f" Trains: {what} (~£{amt})."
    it[f"P{r}"].value = amt
    style(it[f"P{r}"], it[f"G{r}"])
    it[f"K{r}"].value = f"=G{r}+H{r}+I{r}+J{r}+O{r}+P{r}"
it[f"P{total_row}"].value = f"=SUM(P5:P{total_row-1})"; style(it[f"P{total_row}"], it[f"G{total_row}"])
it.column_dimensions["P"].width = 14

# ---- Overview: open up a row after Local transport
AT = next(r for r in range(1, 60) if ov[f"A{r}"].value == "Local transport") + 1
REF = re.compile(r"((?:'[^']+'|[A-Za-z_][\w ]*)!)?(\$?)([A-Z]{1,3})(\$?)(\d+)(?::(\$?)([A-Z]{1,3})(\$?)(\d+))?")
def shift(f):
    def one(m):
        if m.group(1):
            return m.group(0)
        bump = lambda n: str(int(n) + 1 if int(n) >= AT else int(n))
        s = f"{m.group(2)}{m.group(3)}{m.group(4)}{bump(m.group(5))}"
        if m.group(7):
            s += f":{m.group(6)}{m.group(7)}{m.group(8)}{bump(m.group(9))}"
        return s
    return "".join(p if p.startswith('"') else REF.sub(one, p) for p in re.split(r'("[^"]*")', f))

maxr, maxc = ov.max_row, ov.max_column
for r in range(maxr, AT - 1, -1):
    for c in range(1, maxc + 1):
        src, dst = ov.cell(r, c), ov.cell(r + 1, c)
        dst.value = src.value; style(dst, src)
    if ov.row_dimensions[r].height:
        ov.row_dimensions[r + 1].height = ov.row_dimensions[r].height
for row in ov.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.value = shift(c.value)
for c in range(1, maxc + 1):
    style(ov.cell(AT, c), ov.cell(AT - 1, c)); ov.cell(AT, c).value = None
ov[f"A{AT}"].value = LABEL
ov[f"B{AT}"].value = f"=Itinerary!P{total_row}"
ov[f"C{AT}"].value = f"=SUMIFS(Tracker!$E$5:$E$1000,Tracker!$C$5:$C$1000,A{AT})"
ov[f"D{AT}"].value = f"=B{AT}-C{AT}"
ov[f"E{AT}"].value = f"=IF($B$3=0,0,B{AT}/$B$3)"
for row in ov.iter_rows(max_row=AT - 2):
    for c in row:
        if isinstance(c.value, str) and "activities, shopping & misc)" in c.value:
            c.value = c.value.replace("local transport, activities, shopping & misc)", "local transport, trains & intercity, activities, shopping & misc)")

# ---- Tracker category drop-down
for dv in tr.data_validations.dataValidation:
    if dv.formula1 and "Local transport" in dv.formula1 and LABEL not in dv.formula1:
        dv.formula1 = dv.formula1.replace("Local transport,", f"Local transport,{LABEL},")

wb.calculation.fullCalcOnLoad = True
wb.save(XLSX)
print(f"Added {LABEL}: £{sum(a for a, _ in MOVE.values())} moved from Local transport")
